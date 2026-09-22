# -*- coding: utf-8 -*-
import sys
from pathlib import Path
import docx

sys.stdout.reconfigure(encoding='utf-8')
folder = Path('Du An/A Tuan Dan So 2')
out_dir = Path('scripts/temp_lesson_text')
out_dir.mkdir(parents=True, exist_ok=True)

for f in sorted(folder.glob('*.docx')):
    doc = docx.Document(f)
    out_file = out_dir / (f.stem + '.txt')
    with open(out_file, 'w', encoding='utf-8') as out:
        out.write(f"DOCUMENT: {f.name}\n")
        out.write(f"{'='*60}\n\n")
        for i, p in enumerate(doc.paragraphs):
            t = p.text.strip()
            if t:
                out.write(f"{t}\n\n")
        if doc.tables:
            out.write("\n--- TABLES ---\n")
            for t_idx, table in enumerate(doc.tables):
                out.write(f"\n[Table {t_idx+1}]\n")
                for row in table.rows:
                    row_data = [cell.text.strip().replace('\n', ' ') for cell in row.cells]
                    out.write(" | ".join(row_data) + "\n")
    print(f"Dumped: {out_file.name} ({out_file.stat().st_size} bytes)")
