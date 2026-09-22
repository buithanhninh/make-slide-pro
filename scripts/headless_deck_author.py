# -*- coding: utf-8 -*-
"""
headless_deck_author.py
Cross-Platform Headless PowerPoint Authoring Engine using python-pptx.
Renders full 16:9 widescreen presentations with Dark Obsidian and Light Pearl themes,
Bento cards, metric KPI containers, and structured pedagogical typography.
"""

import json
from pathlib import Path
from typing import Any, Dict, List, Optional
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

DARK_TOKENS = {
    "colors": {
        "base": "#060B14",
        "surface": "#0B132B",
        "ink": "#FFFFFF",
        "text_heading": "#FFFFFF",
        "muted": "#CBD5E1",
        "text_muted": "#94A3B8",
        "line": "#CBD5E1",
        "brand": "#0284C7",
        "accent": "#10B981",
    }
}

LIGHT_TOKENS = {
    "colors": {
        "base": "#F8FAFC",
        "surface": "#FFFFFF",
        "ink": "#0F172A",
        "text_heading": "#0F172A",
        "muted": "#334155",
        "text_muted": "#64748B",
        "line": "#475569",
        "brand": "#0284C7",
        "accent": "#10B981",
    }
}

def parse_hex_color(hex_str: str) -> RGBColor:
    clean = hex_str.lstrip("#")
    if len(clean) != 6:
        return RGBColor(255, 255, 255)
    return RGBColor(int(clean[0:2], 16), int(clean[2:4], 16), int(clean[4:6], 16))


class HeadlessDeckAuthor:
    def __init__(self, visible: bool = False, theme: str = "DARK", motion_mode: str = "presenter_click"):
        self.set_theme(theme)
        self.motion_mode = motion_mode.lower()
        self.presentation = None

    def set_theme(self, theme: str):
        self.theme = theme.upper()
        self.tokens = LIGHT_TOKENS if self.theme == "LIGHT" else DARK_TOKENS

    def create_deck(self, blueprints_path: Path, output_pptx: Path, theme: Optional[str] = None) -> Path:
        if theme:
            self.set_theme(theme)

        with open(blueprints_path, "r", encoding="utf-8") as f:
            blueprints = json.load(f)

        output_pptx = Path(output_pptx).resolve()
        output_pptx.parent.mkdir(parents=True, exist_ok=True)
        if output_pptx.exists():
            try:
                output_pptx.unlink()
            except Exception:
                pass

        prs = Presentation()
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        blank_layout = prs.slide_layouts[6]

        base_color = parse_hex_color(self.tokens["colors"]["base"])
        surface_color = parse_hex_color(self.tokens["colors"]["surface"])
        ink_color = parse_hex_color(self.tokens["colors"]["ink"])
        muted_color = parse_hex_color(self.tokens["colors"]["muted"])
        text_muted_color = parse_hex_color(self.tokens["colors"]["text_muted"])
        brand_color = parse_hex_color(self.tokens["colors"]["brand"])

        slides_spec = blueprints.get("slides", [])
        total_slides = len(slides_spec)

        for s_idx, spec in enumerate(slides_spec, start=1):
            slide = prs.slides.add_slide(blank_layout)
            slide.background.fill.solid()
            slide.background.fill.fore_color.rgb = base_color

            role = spec.get("role", "CONTENT").upper()
            title = spec.get("assertion_title", f"Slide {s_idx}")
            claim = spec.get("primary_claim", "")
            section = spec.get("section", "")
            footer = spec.get("source_footer", "Make Slide Pro V8.6.0 - Canonical Enterprise")

            if role == "COVER":
                kicker_box = slide.shapes.add_textbox(Inches(1.2), Inches(1.8), Inches(10.9), Inches(0.5))
                ktf = kicker_box.text_frame
                ktf.word_wrap = True
                kp = ktf.paragraphs[0]
                kp.text = (section or "BÁO CÁO CHIẾN LƯỢC TOÀN DIỆN").upper()
                kp.font.size = Pt(14)
                kp.font.bold = True
                kp.font.color.rgb = brand_color

                title_box = slide.shapes.add_textbox(Inches(1.2), Inches(2.4), Inches(10.9), Inches(2.2))
                ttf = title_box.text_frame
                ttf.word_wrap = True
                tp = ttf.paragraphs[0]
                tp.text = title
                tp.font.size = Pt(36)
                tp.font.bold = True
                tp.font.color.rgb = ink_color

                if claim:
                    claim_box = slide.shapes.add_textbox(Inches(1.2), Inches(4.7), Inches(10.9), Inches(1.5))
                    ctf = claim_box.text_frame
                    ctf.word_wrap = True
                    cp = ctf.paragraphs[0]
                    cp.text = claim
                    cp.font.size = Pt(18)
                    cp.font.color.rgb = muted_color
            else:
                kicker_text = section or f"PHẦN {s_idx:02d}"
                kicker_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.7), Inches(0.35))
                ktf = kicker_box.text_frame
                ktf.word_wrap = True
                kp = ktf.paragraphs[0]
                kp.text = kicker_text.upper()
                kp.font.size = Pt(11)
                kp.font.bold = True
                kp.font.color.rgb = brand_color

                title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11.7), Inches(0.9))
                ttf = title_box.text_frame
                ttf.word_wrap = True
                tp = ttf.paragraphs[0]
                tp.text = title
                tp.font.size = Pt(22)
                tp.font.bold = True
                tp.font.color.rgb = ink_color

                if claim:
                    claim_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.7), Inches(11.7), Inches(0.6))
                    ctf = claim_box.text_frame
                    ctf.word_wrap = True
                    cp = ctf.paragraphs[0]
                    cp.text = claim
                    cp.font.size = Pt(14)
                    cp.font.color.rgb = muted_color

                data_points = spec.get("data_points", [])
                cards = spec.get("cards", [])
                takeaways = spec.get("key_takeaways", []) or spec.get("takeaways", [])

                if data_points:
                    n_dp = min(4, len(data_points))
                    card_w = Inches((11.7 - (n_dp - 1) * 0.3) / n_dp)
                    card_h = Inches(3.8)
                    top_y = Inches(2.5)
                    for dp_idx, dp in enumerate(data_points[:n_dp]):
                        left_x = Inches(0.8) + dp_idx * (card_w + Inches(0.3))
                        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, card_w, card_h)
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = surface_color
                        shape.line.color.rgb = brand_color
                        shape.line.width = Pt(1.5)

                        stf = shape.text_frame
                        stf.word_wrap = True
                        stf.margin_top = Inches(0.3)
                        stf.margin_left = Inches(0.3)

                        val_p = stf.paragraphs[0]
                        val_p.text = str(dp.get("value", ""))
                        val_p.font.size = Pt(36)
                        val_p.font.bold = True
                        val_p.font.color.rgb = brand_color

                        lbl_p = stf.add_paragraph()
                        lbl_p.text = str(dp.get("label", ""))
                        lbl_p.font.size = Pt(14)
                        lbl_p.font.bold = True
                        lbl_p.font.color.rgb = ink_color

                        desc = dp.get("desc", "") or dp.get("growth", "")
                        if desc:
                            desc_p = stf.add_paragraph()
                            desc_p.text = str(desc)
                            desc_p.font.size = Pt(12)
                            desc_p.font.color.rgb = muted_color
                elif cards:
                    n_cards = min(3, len(cards))
                    card_w = Inches((11.7 - (n_cards - 1) * 0.3) / n_cards)
                    card_h = Inches(3.8)
                    top_y = Inches(2.5)
                    for c_idx, c in enumerate(cards[:n_cards]):
                        left_x = Inches(0.8) + c_idx * (card_w + Inches(0.3))
                        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_x, top_y, card_w, card_h)
                        shape.fill.solid()
                        shape.fill.fore_color.rgb = surface_color
                        shape.line.color.rgb = brand_color
                        shape.line.width = Pt(1)

                        stf = shape.text_frame
                        stf.word_wrap = True
                        stf.margin_top = Inches(0.3)
                        stf.margin_left = Inches(0.3)

                        tp_c = stf.paragraphs[0]
                        tp_c.text = c.get("title", f"Mục {c_idx+1}")
                        tp_c.font.size = Pt(16)
                        tp_c.font.bold = True
                        tp_c.font.color.rgb = ink_color

                        c_desc = c.get("desc", "") or c.get("text", "") or c.get("content", "")
                        if c_desc:
                            dp_c = stf.add_paragraph()
                            dp_c.text = c_desc
                            dp_c.font.size = Pt(13)
                            dp_c.font.color.rgb = muted_color
                else:
                    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.5), Inches(11.7), Inches(3.8))
                    shape.fill.solid()
                    shape.fill.fore_color.rgb = surface_color
                    shape.line.color.rgb = brand_color
                    shape.line.width = Pt(1)

                    stf = shape.text_frame
                    stf.word_wrap = True
                    stf.margin_top = Inches(0.4)
                    stf.margin_left = Inches(0.4)

                    items = takeaways or spec.get("bullet_points", []) or [claim]
                    for idx, item in enumerate(items[:6]):
                        p = stf.paragraphs[0] if idx == 0 else stf.add_paragraph()
                        p.text = f"•  {item}"
                        p.font.size = Pt(14)
                        p.font.color.rgb = ink_color

            footer_box = slide.shapes.add_textbox(Inches(0.8), Inches(6.9), Inches(10), Inches(0.3))
            ftf = footer_box.text_frame
            fp = ftf.paragraphs[0]
            fp.text = footer
            fp.font.size = Pt(10)
            fp.font.color.rgb = text_muted_color

            num_box = slide.shapes.add_textbox(Inches(11.5), Inches(6.9), Inches(1.0), Inches(0.3))
            ntf = num_box.text_frame
            np = ntf.paragraphs[0]
            np.text = f"{s_idx}/{total_slides}"
            np.font.size = Pt(10)
            np.alignment = PP_ALIGN.RIGHT
            np.font.color.rgb = text_muted_color

        prs.save(str(output_pptx))
        return output_pptx

    def close(self):
        pass
