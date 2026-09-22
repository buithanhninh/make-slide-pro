# -*- coding: utf-8 -*-
import sys
from pathlib import Path
import docx

sys.stdout.reconfigure(encoding='utf-8')
folder = Path('Du An/A Tuan Dan So 2')

for f in sorted(folder.glob('*.docx')):
    doc = docx.Document(f)
    print(f"\n{'='*70}\nFILE: {f.name}\n{'='*70}")
    print(f"Total paragraphs: {len(doc.paragraphs)}, Total tables: {len(doc.tables)}")
    
    # Extract structural outline
    print("--- STRUCTURAL OUTLINE ---")
    count = 0
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue
        if any(txt.startswith(prefix) for prefix in ['BÀI', 'Bài', 'I.', 'II.', 'III.', 'IV.', 'V.', 'VI.', '1.', '2.', '3.', '4.', '5.', '6.', '7.', '8.', '9.', '10.']):
            if len(txt) < 120:
                print(f"  {txt}")
                count += 1
                if count >= 30:
                    print("  ... (more headings)")
                    break
